from docx import Document
from docx.shared import Pt

class CreateDoc:

    def __init__(self):
        super().__init__()

    doc = Document()

    def create(self, name_file_path):

        self.doc.add_heading('Hi class base docx')
        parag = self.doc.add_paragraph('Lorem ipsum dolor sit amet consectetur, adipisicing elit. Sint eos dignissimos, voluptate pariatur dolores quidem! Aperiam cupiditate illum deleniti? Quidem quos nostrum, doloribus consequuntur temporibus blanditiis maiores beatae pariatur nihil.')
        parag.add_run('i will use ')
        parag.add_run('blod words ').bold = True
        parag.add_run('i will use ')
        parag.add_run('italic words ').italic = True

        text_font = self.doc.add_paragraph('first text ').add_run('test text for font!')
        font = text_font.font
        font.name = 'system'
        font.size = Pt(22)

        self.doc.save(name_file_path)

    def read_docx(self, file_name):

        doc = Document(file_name)

        parag = doc.paragraphs

        for i in parag:
            print(i.text)
    
    def replace_text(self, file_name, first_text, second_text, file_name_new):

        doc = Document(file_name)

        for i in doc.paragraphs:
            i.text = i.text.replace(first_text, second_text)

        doc.save(file_name_new)


if __name__ == "__main__":
    main = CreateDoc()
    # main.create('doc.docx')
    # main.read_docx('doc.docx')
    # main.replace_text('doc.docx', 'Satrap', 'Satrap18', 'doc.docx')
    